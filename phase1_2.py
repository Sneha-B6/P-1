import streamlit as st
from langchain_ollama import OllamaLLM
import os
import pdfplumber
from fpdf import FPDF
import re
import base64
from docx import Document
import nltk

# Initialize required nltk resources
nltk.download('wordnet')

# Backend Module: Extraction and Generation
class BRDProcessor:
    def __init__(self):
        self.llm = OllamaLLM(model="llama3.2")

    def extract_text(self, uploaded_file, file_type):
        if file_type == "pdf":
            return self._extract_text_from_pdf(uploaded_file)
        elif file_type == "docx":
            return self._extract_text_from_docx(uploaded_file)
        elif file_type == "txt":
            return self._extract_text_from_txt(uploaded_file)
        return None

    def _extract_text_from_pdf(self, pdf_file):
        with pdfplumber.open(pdf_file) as pdf:
            text = ''.join([page.extract_text() + '\n' for page in pdf.pages])
        return text

    def _extract_text_from_docx(self, docx_file):
        doc = Document(docx_file)
        return '\n'.join([paragraph.text for paragraph in doc.paragraphs])

    def _extract_text_from_txt(self, txt_file):
        return txt_file.read().decode('utf-8')

    def create_user_story(self, brd_content, prompt=""):
        user_story_prompt = (
            f"Using the provided BRD content, generate a user story in a structured format. "
            f"Using the provided user_prompt, restructure the user story and if no user_prompt is given do not take it into consideration. "
            f"The content involves cargo management, operational processes, and task flows. "
            f"Output the user story with the following format:\n"
            f"Actors: (List of all actors)\n"
            f"Preconditions: (Conditions required before starting)\n"
            f"Main Flow: (Detailed steps)\n"
            f"Postconditions: (Expected outcomes)\n"
            f"Exceptions: (Potential deviations)\n\n"
            f"Content:\n{brd_content}\nuser_prompt:\n{prompt}"
        )
        return self.llm.invoke(user_story_prompt)

    def create_use_case(self, user_story):
        use_case_prompt = (
            f"Extract the following information from the generated user story:\n"
            f"Actors: (List of all actors)\n"
            f"Preconditions: (Conditions required before starting)\n"
            f"Main Flow: (Detailed steps)\n"
            f"Postconditions: (Expected outcomes)\n"
            f"Exceptions: (Potential deviations)\n\n"
            f"User Story:\n{user_story}"
        )
        use_case = self.llm.invoke(use_case_prompt)
        return self._extract_use_case_info(use_case)

    def _extract_use_case_info(self, user_story):
        patterns = {
            "Actors": re.compile(r"Actors\s*:\s*(.*?)(?=Preconditions|$)", re.DOTALL | re.IGNORECASE),
            "Preconditions": re.compile(r"Preconditions\s*:\s*(.*?)(?=Main Flow|$)", re.DOTALL | re.IGNORECASE),
            "Main Flow": re.compile(r"Main Flow\s*:\s*(.*?)(?=Postconditions|$)", re.DOTALL | re.IGNORECASE),
            "Postconditions": re.compile(r"Postconditions\s*:\s*(.*?)(?=Exceptions|$)", re.DOTALL | re.IGNORECASE),
            "Exceptions": re.compile(r"Exceptions\s*:\s*(.*)", re.DOTALL | re.IGNORECASE)
        }
        use_case_info = {}
        for key, pattern in patterns.items():
            match = pattern.search(user_story)
            use_case_info[key] = match.group(1).strip() if match else "No information available."
        return use_case_info

    def generate_pdf(self, use_case, file_name):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_draw_color(0, 0, 0)
        pdf.set_line_width(0.5)
        pdf.rect(5, 5, 200, 287)
        pdf.set_font("Times", 'B', size=16)
        pdf.cell(0, 10, "Use Case Document", ln=True, align="C")
        pdf.ln(10)
        pdf.set_font("Times", 'B', size=12)
        for key, value in use_case.items():
            pdf.cell(0, 10, f"{key}:", ln=True)
            pdf.set_font("Times", size=12)
            content_words = value.split()
            current_text = ""
            for word in content_words:
                if len(current_text) + len(word) + 1 > 50:
                    pdf.cell(0, 10, current_text, ln=True)
                    current_text = word
                    if pdf.get_y() + 10 > pdf.h - 20:
                        pdf.add_page()
                        pdf.set_draw_color(0, 0, 0)
                        pdf.set_line_width(0.5)
                        pdf.rect(5, 5, 200, 287)
                else:
                    current_text += (" " + word) if current_text else word
            if current_text:
                pdf.cell(0, 10, current_text, ln=True)
            pdf.set_font("Times", 'B', size=12)
            pdf.cell(0, 5, '', ln=True)
        pdf.output(file_name)


# UI Module: Display and Download
class UseCaseApp:
    def __init__(self):
        self.processor = BRDProcessor()

    def run(self):
        st.set_page_config(page_title="BRD to Use Case Generator", layout="wide")
        self.set_background("/Users/snehamariarejo/Documents/2.png")

        st.title("BRD to Use Case Generator")
        st.write("Easily convert your BRD into structured user stories and use cases for effective cargo management.")
        
        self.sidebar_instructions()
        uploaded_file, prompt = self.get_user_inputs()

        if uploaded_file:
            file_type = uploaded_file.name.split('.')[-1]
            extracted_text = self.processor.extract_text(uploaded_file, file_type)

            if extracted_text:
                st.success("BRD file uploaded and text extracted successfully.")
                st.write("### Extracted Text:")
                st.write(extracted_text[:500] + "...")

                if st.button("Generate User Story"):
                    self.generate_outputs(extracted_text, prompt)
            else:
                st.error("Failed to extract text from the uploaded file. Please check the file format and content.")

    def set_background(self, image_path):
        image_base64 = self.convert_image_to_base64(image_path)
        st.markdown(
            f"""
            <style>
                .stApp {{
                    background-image: url("data:image/jpeg;base64,{image_base64}");
                    background-repeat: no-repeat;
                    background-size: cover;
                    background-attachment: fixed;
                }}
            </style>
            """, unsafe_allow_html=True
        )

    def sidebar_instructions(self):
        with st.sidebar:
            st.header("📄 Instructions")
            st.markdown("<ul><li>Upload a BRD in PDF, DOCX, or TXT format.</li><li>Optionally, specify a focus area.</li><li>Download the generated use case as a PDF.</li></ul>", unsafe_allow_html=True)

    def get_user_inputs(self):
        uploaded_file = st.file_uploader("📄 Upload a BRD file:", type=["pdf", "docx", "txt"])
        prompt = st.text_input("📝 Optional prompt for use case generation (specify region):")
        return uploaded_file, prompt

    def generate_outputs(self, extracted_text, prompt):
        user_story = self.processor.create_user_story(extracted_text, prompt)
        
        col1, col2 = st.columns(2)
        with col1:
            with st.expander("USER STORY"):
                st.subheader("Generated User Story")
                st.write(user_story)
        
        with col2:
            with st.expander("USE CASE"):
                use_case = self.processor.create_use_case(user_story)
                self.display_use_case(use_case)

        # PDF download link
        use_case_pdf_name = "use_case.pdf"
        self.processor.generate_pdf(use_case, use_case_pdf_name)
        with open(use_case_pdf_name, "rb") as pdf_file:
            st.download_button("📥 Download Use Case PDF", pdf_file, file_name=use_case_pdf_name)

    def display_use_case(self, use_case):
        st.subheader("Generated Use Case")
        for key, value in use_case.items():
            st.markdown(f"{key}:")
            st.markdown(value)
            st.markdown("")

    @staticmethod
    def convert_image_to_base64(image_path):
        with open(image_path, "rb") as img_file:
            return base64.b64encode(img_file.read()).decode("utf-8")


if __name__ == "__main__":
    app = UseCaseApp()
    app.run()